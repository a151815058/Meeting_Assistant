import base64
import email
import io
from datetime import datetime, timezone
from email import policy
from urllib.parse import unquote

import pytest
from docx import Document
from docx.oxml.ns import qn

from app.extensions import db as _db
from app.minutes import export, generator
from app.models.audit import AuditLog
from app.models.meeting import Meeting, Participant
from app.models.template import Minutes
from app.models.user import User
from app.notifications import mailer
from app.notifications.senders import Attachment, GraphSender, MailSender, OutgoingMail, Recipient, SendResult

MD = """# Q3 預算會議 會議記錄

## 決議
1. 行銷預算**增加兩成**
2. 下週五前提交
   - 子項目 A

| 負責人 | 事項 |
|---|---|
| 李小姐 | 預算表 |

> 備註：<script>alert(1)</script>

參考 [網站](https://example.com)

```
code line
```
---
"""


class FakeSender(MailSender):
    backend = "fake"

    def __init__(self):
        self.sent = []

    def send(self, mail):
        self.sent.append(mail)
        return SendResult(backend=self.backend, message_id="m1")


@pytest.fixture(autouse=True)
def sender(app):
    mailer.reset_state()
    generator.reset_jobs()
    fake = FakeSender()
    app.extensions["mail_sender"] = fake
    yield fake
    mailer.reset_state()


def _login(client, app, email="org@example.com"):
    with app.app_context():
        user = User(email=email, display_name="王經理")
        _db.session.add(user)
        _db.session.commit()
        user_id = user.id
    with client.session_transaction() as sess:
        sess["_user_id"] = user_id
        sess["_fresh"] = True
    return user_id


def _meeting(app, user_id, *, title="Q3 預算會議", minutes=True):
    with app.app_context():
        meeting = Meeting(organizer_id=user_id, title=title, status="minutes_ready",
                          scheduled_start=datetime(2026, 9, 25, 6, tzinfo=timezone.utc),
                          scheduled_end=datetime(2026, 9, 25, 7, tzinfo=timezone.utc))
        _db.session.add(meeting)
        _db.session.flush()
        _db.session.add(Participant(meeting_id=meeting.id, email="li@example.com", display_name="李小姐"))
        if minutes:
            _db.session.add(Minutes(meeting_id=meeting.id, content_markdown=MD, llm_provider="fake", llm_model="m"))
        _db.session.commit()
        return meeting.id


def _docx_text(data: bytes) -> tuple[Document, str]:
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    text += "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    return doc, text


# --- TC-41: Markdown parsing shared by both formats --------------------------------------------

def test_markdown_is_parsed_into_blocks():
    blocks = export.parse_markdown(MD)

    kinds = [b.kind for b in blocks]
    assert kinds == ["heading", "heading", "list", "table", "quote", "para", "code", "hr"]
    lst = blocks[2]
    assert lst.ordered and len(lst.items) == 2
    assert [r.text for r in lst.items[0][0].runs if r.bold] == ["增加兩成"]
    assert lst.items[1][1].kind == "list" and not lst.items[1][1].ordered  # nested bullet list
    assert blocks[3].header_rows == 1 and [r.text for r in blocks[3].rows[1][0]] == ["李小姐"]
    quote_text = "".join(r.text for r in blocks[4].items[0][0].runs)
    assert "<script>alert(1)</script>" in quote_text  # raw HTML stays literal text
    assert "".join(r.text for r in blocks[5].runs) == "參考 網站（https://example.com）"


def test_filename_is_sanitised():
    meeting = type("M", (), {"title": 'a/b:c*?"<>|\\d\r\n', "scheduled_start": datetime(2026, 9, 25, 6, tzinfo=timezone.utc),
                             "minutes": None})()
    from flask import Flask

    app = Flask(__name__)
    app.config["DISPLAY_TIMEZONE"] = "Asia/Taipei"
    with app.app_context():
        assert export.filename(meeting, "pdf") == "會議記錄_a_b_c_d_20260925.pdf"


# --- TC-41: download --------------------------------------------------------------------------

def test_download_word(client, app):
    user_id = _login(client, app)
    meeting_id = _meeting(app, user_id)

    resp = client.get(f"/meetings/{meeting_id}/minutes/export/docx")

    assert resp.status_code == 200 and resp.mimetype == export.FORMATS["docx"]
    assert resp.headers["Cache-Control"] == "no-store"
    disposition = resp.headers["Content-Disposition"]
    assert "attachment" in disposition and "會議記錄_Q3 預算會議_20260925.docx" in unquote(disposition)
    doc, text = _docx_text(resp.data)
    assert "會議記錄：Q3 預算會議" in text and "2026-09-25（五）14:00–15:00" in text and "李小姐" in text
    assert "1. 行銷預算增加兩成" in text and "• 子項目 A" in text and "預算表" in text
    assert "AI 依會議逐字稿產生" in text
    assert doc.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == export.DOCX_FONT
    with app.app_context():
        audit = AuditLog.query.filter_by(action="minutes.exported").one()
        assert audit.event_metadata["format"] == "docx"


def test_download_pdf(client, app):
    user_id = _login(client, app)
    meeting_id = _meeting(app, user_id)

    resp = client.get(f"/meetings/{meeting_id}/minutes/export/pdf")

    assert resp.status_code == 200 and resp.mimetype == "application/pdf"
    assert resp.data.startswith(b"%PDF") and b"/FontFile2" in resp.data  # CJK font embedded (subset)


def test_download_reflects_saved_edits(client, app):
    user_id = _login(client, app)
    meeting_id = _meeting(app, user_id)
    client.post(f"/meetings/{meeting_id}/minutes", data={"content_markdown": "# 修改後的版本"})

    _, text = _docx_text(client.get(f"/meetings/{meeting_id}/minutes/export/docx").data)
    assert "修改後的版本" in text and "增加兩成" not in text


def test_export_access_control(client, app):
    owner = _login(client, app, email="owner@example.com")
    theirs = _meeting(app, owner)
    me = _login(client, app, email="me@example.com")
    no_minutes = _meeting(app, me, minutes=False)
    mine = _meeting(app, me)

    assert client.get(f"/meetings/{theirs}/minutes/export/docx").status_code == 404
    assert client.get(f"/meetings/{no_minutes}/minutes/export/pdf").status_code == 404
    assert client.get(f"/meetings/{mine}/minutes/export/html").status_code == 404


def test_pdf_without_cjk_font_shows_message(client, app):
    app.config["PDF_FONT_PATH"] = r"C:\no\such\font.ttf"
    user_id = _login(client, app)
    meeting_id = _meeting(app, user_id)

    resp = client.get(f"/meetings/{meeting_id}/minutes/export/pdf", follow_redirects=True)

    assert "找不到中文字型" in resp.get_data(as_text=True)
    with app.app_context():
        assert AuditLog.query.filter_by(action="minutes.exported").count() == 0


def test_minutes_page_offers_downloads_and_attachments(client, app):
    user_id = _login(client, app)
    meeting_id = _meeting(app, user_id)

    page = client.get(f"/meetings/{meeting_id}/minutes").get_data(as_text=True)
    assert "下載 Word" in page and "下載 PDF" in page
    assert 'name="attach" value="docx"' in page and 'name="attach" value="pdf"' in page


# --- TC-41: attachments -----------------------------------------------------------------------

def test_send_with_word_and_pdf_attachments(client, app, sender):
    user_id = _login(client, app)
    meeting_id = _meeting(app, user_id)

    client.post(f"/meetings/{meeting_id}/minutes/send", data={"attach": ["pdf", "docx", "exe"]})

    mail = sender.sent[0]
    assert [(a.filename, a.mimetype) for a in mail.attachments] == [
        ("會議記錄_Q3 預算會議_20260925.docx", export.FORMATS["docx"]),
        ("會議記錄_Q3 預算會議_20260925.pdf", "application/pdf"),
    ]
    assert mail.attachments[1].data.startswith(b"%PDF")
    with app.app_context():
        assert AuditLog.query.filter_by(action="mail.sent").one().event_metadata["attachments"] == ["docx", "pdf"]


def test_send_without_attachments_by_default(client, app, sender):
    user_id = _login(client, app)
    meeting_id = _meeting(app, user_id)
    client.post(f"/meetings/{meeting_id}/minutes/send")
    assert sender.sent[0].attachments == ()


@pytest.mark.parametrize("config,message", [
    ({"MAIL_MAX_ATTACHMENT_BYTES": 100}, "附件太大"),
    ({"PDF_FONT_PATH": r"C:\no\such\font.ttf"}, "找不到中文字型"),
])
def test_attachment_problems_block_sending(client, app, sender, config, message):
    app.config.update(config)
    user_id = _login(client, app)
    meeting_id = _meeting(app, user_id)

    resp = client.post(f"/meetings/{meeting_id}/minutes/send", data={"attach": ["pdf"]}, follow_redirects=True)

    assert message in resp.get_data(as_text=True) and sender.sent == []
    with app.app_context():
        assert _db.session.get(Meeting, meeting_id).minutes.status == "draft"


MAIL = OutgoingMail(sender=Recipient("org@example.com", "王經理"), recipients=(Recipient("li@example.com", "李小姐"),),
                    subject="s", body="b",
                    attachments=(Attachment("會議記錄_預算.pdf", "application/pdf", b"%PDF-1.4 test"),))


def test_mime_attachment_keeps_chinese_filename():
    msg = email.message_from_bytes(MAIL.to_mime().as_bytes(), policy=policy.default)

    part, = list(msg.iter_attachments())
    assert part.get_filename() == "會議記錄_預算.pdf" and part.get_content_type() == "application/pdf"
    assert part.get_content() == b"%PDF-1.4 test"


def test_graph_payload_includes_attachments(mocker):
    post = mocker.patch("app.notifications.senders.requests.post",
                        return_value=type("R", (), {"status_code": 202, "text": ""})())

    GraphSender("tok").send(MAIL)

    att, = post.call_args.kwargs["json"]["message"]["attachments"]
    assert att["@odata.type"] == "#microsoft.graph.fileAttachment" and att["name"] == "會議記錄_預算.pdf"
    assert base64.b64decode(att["contentBytes"]) == b"%PDF-1.4 test"
