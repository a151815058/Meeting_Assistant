"""Export minutes as Word (.docx) or PDF (REQ-41).

The saved Markdown is parsed once (markdown-it, raw HTML disabled) into a small block model,
which both renderers walk, so the two formats have the same structure. The document starts
with the meeting details (title, time, organizer, participants), then the minutes.
"""
import io
import logging
import os
import re
from dataclasses import dataclass, field

from markdown_it import MarkdownIt

from app.filters import localtime

# fontTools logs a warning for every OpenType table it cannot subset; they are harmless.
logging.getLogger("fontTools").setLevel(logging.ERROR)

FORMATS = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
FORMAT_LABELS = {"docx": "Word", "pdf": "PDF"}

# Tried in order when PDF_FONT_PATH is not set: Windows (Microsoft JhengHei), then the
# Debian/Ubuntu fonts-noto-cjk package used by the Dockerfile.
_FONT_CANDIDATES = [
    (r"C:\Windows\Fonts\msjh.ttc", r"C:\Windows\Fonts\msjhbd.ttc"),
    ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
]
DOCX_FONT = "Microsoft JhengHei"


class ExportError(Exception):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


# --- Markdown -> block model -----------------------------------------------------------------

@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False


@dataclass
class Block:
    kind: str  # heading | para | list | table | code | hr | quote
    level: int = 0
    runs: list[Run] = field(default_factory=list)
    ordered: bool = False
    start: int = 1
    items: list[list["Block"]] = field(default_factory=list)  # list items / quote children ([0])
    rows: list[list[list[Run]]] = field(default_factory=list)  # table rows -> cells -> runs
    header_rows: int = 0
    text: str = ""


_md = MarkdownIt("commonmark", {"html": False}).enable(["table", "strikethrough"])


def _inline_runs(token) -> list[Run]:
    runs, bold, italic, href = [], 0, 0, None
    for child in token.children or []:
        t = child.type
        if t == "text":
            runs.append(Run(child.content, bold > 0, italic > 0))
        elif t == "code_inline":
            runs.append(Run(child.content, bold > 0, italic > 0, code=True))
        elif t in ("softbreak", "hardbreak"):
            runs.append(Run("\n"))
        elif t == "strong_open":
            bold += 1
        elif t == "strong_close":
            bold -= 1
        elif t == "em_open":
            italic += 1
        elif t == "em_close":
            italic -= 1
        elif t == "link_open":
            href = child.attrGet("href")
        elif t == "link_close":
            if href and not any(r.text == href for r in runs[-1:]):
                runs.append(Run(f"（{href}）"))
            href = None
        elif t == "image":
            runs.append(Run(child.content or "[圖片]"))
    return runs


def _parse_blocks(tokens, i: int, stop: str | None) -> tuple[list[Block], int]:
    blocks = []
    while i < len(tokens):
        tok = tokens[i]
        if stop and tok.type == stop:
            return blocks, i + 1
        if tok.type == "heading_open":
            blocks.append(Block("heading", level=int(tok.tag[1]), runs=_inline_runs(tokens[i + 1])))
            i += 3
        elif tok.type == "paragraph_open":
            blocks.append(Block("para", runs=_inline_runs(tokens[i + 1])))
            i += 3
        elif tok.type in ("bullet_list_open", "ordered_list_open"):
            close = tok.type.replace("_open", "_close")
            block = Block("list", ordered=tok.type == "ordered_list_open", start=int(tok.attrGet("start") or 1))
            i += 1
            while tokens[i].type != close:
                item, i = _parse_blocks(tokens, i + 1, "list_item_close")  # skip list_item_open
                block.items.append(item)
            blocks.append(block)
            i += 1
        elif tok.type == "blockquote_open":
            children, i = _parse_blocks(tokens, i + 1, "blockquote_close")
            blocks.append(Block("quote", items=[children]))
        elif tok.type == "table_open":
            block, i = _parse_table(tokens, i + 1)
            blocks.append(block)
        elif tok.type in ("fence", "code_block"):
            blocks.append(Block("code", text=tok.content.rstrip("\n")))
            i += 1
        elif tok.type == "hr":
            blocks.append(Block("hr"))
            i += 1
        else:
            i += 1
    return blocks, i


def _parse_table(tokens, i: int) -> tuple[Block, int]:
    block, in_head, row = Block("table"), False, None
    while tokens[i].type != "table_close":
        t = tokens[i].type
        if t == "thead_open":
            in_head = True
        elif t == "thead_close":
            in_head = False
        elif t == "tr_open":
            row = []
        elif t == "tr_close":
            block.rows.append(row)
            block.header_rows += 1 if in_head else 0
        elif t == "inline":
            row.append(_inline_runs(tokens[i]))
        i += 1
    return block, i + 1


def parse_markdown(text: str) -> list[Block]:
    return _parse_blocks(_md.parse(text), 0, None)[0]


# --- shared document header -------------------------------------------------------------------

def _meeting_info(meeting) -> list[tuple[str, str]]:
    info = []
    if meeting.scheduled_start:
        when = localtime(meeting.scheduled_start, "%Y-%m-%d（%W）%H:%M")
        if meeting.scheduled_end:
            when += "–" + localtime(meeting.scheduled_end, "%H:%M")
        info.append(("會議時間", when))
    organizer = meeting.organizer
    info.append(("主辦人", organizer.display_name or organizer.email))
    names = [p.display_name or p.email for p in sorted(meeting.participants, key=lambda p: p.email.lower())]
    info.append(("與會者", "、".join(names) if names else "—"))
    return info


FOOTNOTE = "本會議記錄由 AI 依會議逐字稿產生，並經主辦人確認。"


def filename(meeting, fmt: str) -> str:
    title = re.sub(r'[\\/:*?"<>|\x00-\x1f\x7f]+', "_", meeting.title).strip(" ._")[:80] or "會議"
    stamp = localtime(meeting.scheduled_start or meeting.minutes.generated_at, "%Y%m%d")
    return f"會議記錄_{title}_{stamp}.{fmt}"


def render(meeting, fmt: str, config) -> bytes:
    if fmt == "docx":
        return render_docx(meeting)
    if fmt == "pdf":
        return render_pdf(meeting, config)
    raise ExportError("unknown_format")


# --- Word -----------------------------------------------------------------------------------

def _set_style_font(style, name: str) -> None:
    from docx.oxml.ns import qn

    style.font.name = name
    # CJK text uses the East Asian font slot, which font.name does not set.
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def _docx_runs(paragraph, runs: list[Run]) -> None:
    for r in runs:
        run = paragraph.add_run(r.text)
        run.bold, run.italic = r.bold or None, r.italic or None
        if r.code:
            run.font.name = "Consolas"


def _docx_blocks(doc, blocks: list[Block], depth: int = 0) -> None:
    from docx.shared import Pt

    for b in blocks:
        if b.kind == "heading":
            _docx_runs(doc.add_heading(level=min(b.level, 4)), b.runs)
        elif b.kind == "para":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18 * depth) if depth else None
            _docx_runs(p, b.runs)
        elif b.kind == "list":
            # Numbers/bullets are written as text: Word's automatic numbering would continue
            # across separate lists in python-docx.
            for n, item in enumerate(b.items, start=b.start):
                marker = f"{n}. " if b.ordered else "• "
                first, rest = (item[0], item[1:]) if item and item[0].kind == "para" else (Block("para"), item)
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18 * (depth + 1))
                p.paragraph_format.first_line_indent = Pt(-12)
                p.add_run(marker)
                _docx_runs(p, first.runs)
                _docx_blocks(doc, rest, depth + 1)
        elif b.kind == "table" and b.rows:
            cols = max(len(r) for r in b.rows)
            table = doc.add_table(rows=len(b.rows), cols=cols, style="Table Grid")
            for ri, row in enumerate(b.rows):
                for ci, cell_runs in enumerate(row):
                    p = table.cell(ri, ci).paragraphs[0]
                    _docx_runs(p, [Run(r.text, r.bold or ri < b.header_rows, r.italic, r.code) for r in cell_runs])
            doc.add_paragraph()
        elif b.kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        elif b.kind == "hr":
            doc.add_paragraph("—" * 20)
        elif b.kind == "quote":
            for child in b.items[0]:
                _docx_blocks(doc, [child], depth + 1)


def render_docx(meeting) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        _set_style_font(doc.styles[name], DOCX_FONT)
    doc.core_properties.title = f"會議記錄：{meeting.title}"
    doc.core_properties.author = meeting.organizer.display_name or meeting.organizer.email

    doc.add_heading(f"會議記錄：{meeting.title}", level=0)
    info = doc.add_table(rows=0, cols=2, style="Table Grid")
    for label, value in _meeting_info(meeting):
        cells = info.add_row().cells
        cells[0].paragraphs[0].add_run(label).bold = True
        cells[1].paragraphs[0].add_run(value)
    doc.add_paragraph()

    _docx_blocks(doc, parse_markdown(meeting.minutes.content_markdown))

    note = doc.add_paragraph().add_run(FOOTNOTE)
    note.font.size = Pt(9)
    note.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --- PDF ------------------------------------------------------------------------------------

def _font_paths(config) -> tuple[str, str]:
    regular = config.get("PDF_FONT_PATH") or ""
    if regular:
        if not os.path.exists(regular):
            raise ExportError("pdf_font_missing")
        bold = config.get("PDF_BOLD_FONT_PATH") or regular
        return regular, bold if os.path.exists(bold) else regular
    for reg, bold in _FONT_CANDIDATES:
        if os.path.exists(reg):
            return reg, bold if os.path.exists(bold) else reg
    raise ExportError("pdf_font_missing")


_HEADING_SIZES = {1: 17, 2: 14.5, 3: 12.5}
_BODY, _LINE = 10.5, 6.2  # font size (pt), line height (mm)


def render_pdf(meeting, config) -> bytes:
    from fpdf import FPDF, FontFace
    from fpdf.enums import WrapMode, XPos, YPos

    regular, bold = _font_paths(config)

    class _Doc(FPDF):
        def footer(self):
            self.set_y(-12)
            self.set_font("CJK", "", 8)
            self.set_text_color(120)
            self.cell(0, 6, f"{self.page_no()} / {{nb}}", align="C")
            self.set_text_color(0)

    pdf = _Doc(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_font("CJK", "", regular)
    pdf.add_font("CJK", "B", bold)
    pdf.set_title(f"會議記錄：{meeting.title}")
    pdf.set_author(meeting.organizer.display_name or meeting.organizer.email)
    pdf.set_creator("Meeting Assistant")
    pdf.add_page()
    base_margin = pdf.l_margin

    def write_runs(runs, size=_BODY, force_bold=False, line=_LINE):
        for r in runs:
            pdf.set_font("CJK", "B" if (r.bold or force_bold) else "", size)
            pdf.write(line, r.text, wrapmode=WrapMode.CHAR)
        pdf.ln(line)

    def blocks(items, depth=0):
        for b in items:
            pdf.set_left_margin(base_margin + 6 * depth)
            pdf.set_x(pdf.l_margin)
            if b.kind == "heading":
                size = _HEADING_SIZES.get(b.level, _BODY + 0.5)
                pdf.ln(2)
                write_runs(b.runs, size=size, force_bold=True, line=size * 0.5)
                pdf.ln(1)
            elif b.kind == "para":
                write_runs(b.runs)
                pdf.ln(1.5)
            elif b.kind == "list":
                for n, item in enumerate(b.items, start=b.start):
                    pdf.set_left_margin(base_margin + 6 * (depth + 1))
                    pdf.set_x(pdf.l_margin)
                    first, rest = (item[0], item[1:]) if item and item[0].kind == "para" else (Block("para"), item)
                    write_runs([Run(f"{n}. " if b.ordered else "• ")] + first.runs)
                    blocks(rest, depth + 1)
                pdf.ln(1.5)
            elif b.kind == "table" and b.rows:
                pdf.set_font("CJK", "", _BODY - 0.5)
                cols = max(len(r) for r in b.rows)
                with pdf.table(headings_style=FontFace(emphasis="BOLD"), first_row_as_headings=b.header_rows > 0,
                               line_height=_LINE, wrapmode=WrapMode.CHAR) as table:
                    for row in b.rows:
                        cells = ["".join(r.text for r in cell) for cell in row]
                        table.row(cells + [""] * (cols - len(cells)))
                pdf.ln(2)
            elif b.kind == "code":
                pdf.set_font("CJK", "", _BODY - 1)
                pdf.multi_cell(0, _LINE - 0.8, b.text, wrapmode=WrapMode.CHAR, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.ln(1.5)
            elif b.kind == "hr":
                y = pdf.get_y() + 1.5
                pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
                pdf.ln(4)
            elif b.kind == "quote":
                blocks(b.items[0], depth + 1)
        pdf.set_left_margin(base_margin)

    pdf.set_font("CJK", "B", 18)
    pdf.multi_cell(0, 9, f"會議記錄：{meeting.title}", wrapmode=WrapMode.CHAR, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)
    pdf.set_font("CJK", "", _BODY)
    with pdf.table(col_widths=(22, 78), first_row_as_headings=False, line_height=_LINE,
                   wrapmode=WrapMode.CHAR) as table:
        for label, value in _meeting_info(meeting):
            row = table.row()
            row.cell(label, style=FontFace(emphasis="BOLD"))
            row.cell(value)
    pdf.ln(4)

    blocks(parse_markdown(meeting.minutes.content_markdown))

    pdf.ln(3)
    pdf.set_font("CJK", "", 8.5)
    pdf.set_text_color(110)
    pdf.multi_cell(0, 5, FOOTNOTE, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    return bytes(pdf.output())
